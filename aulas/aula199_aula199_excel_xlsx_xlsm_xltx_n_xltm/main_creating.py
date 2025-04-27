# openpyxl - to work with Excel files
# xlsx - Excel Workbook
# xlsm - Excel Macro-Enabled Workbook
# xltx - Excel Template Workbook
# xltm - Excel Macro-Enabled Template Workbook

# With openpyxl, we can create and manipulate Excel files in Python.
# It allows us to read, write, and modify Excel files in the xlsx and xlsm formats.
# Its usefull to create reports, data analysis, and automate tasks in Excel.
# It is a powerful library that provides a lot of functionality for working with Excel files.
# Documentation: https://openpyxl.readthedocs.io/en/stable/index.html


from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

root_folder = Path(__file__).parent
workbook_path = root_folder / 'workbook.xlsx'

workbook = Workbook()  # Create a new workbook
# worksheet: Worksheet = workbook.active  # Get the active worksheet

# Name to the spreadsheet
sheet_name = 'Students'
# Create the spreasheet with the name 'Students'
workbook.create_sheet(sheet_name, 0)
# Get the worksheet by name
worksheet: Worksheet = workbook[sheet_name]

# Remove the default sheet created by openpyxl
if 'Sheet' in workbook.sheetnames:
    std = workbook['Sheet']
    workbook.remove(std)  # Remove the default sheet


# Creating the header
worksheet.cell(1, 1, 'Name')  # Column A
worksheet.cell(1, 2, 'Age')  # Column B
worksheet.cell(1, 3, 'Grade') # Column C


students = [
    ['Ailton', 27,  3],
    ['Maria',  52,  10],
    ['José',   55,  5],
    ['Sti',    35,  8],
    ['Ellen',  34,  3],
]
# print(workbook.sheetnames)  # Print the names of the sheets in the workbook

# for i in range(2, 10):
#     for j in range(1, 4):
#         print(f'Row {i}, Column {j}')

# for i, student_row in enumerate(students, start=2):  # Start from row 2
#     for j, student_column in enumerate(student_row, start=1):
#         worksheet.cell(i, j, student_column)
        

for student in students:
    worksheet.append(student)  # Append the student data to the worksheet

workbook.save(workbook_path)  # Save the workbook to a file
